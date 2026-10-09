import os
from datetime import datetime
from zoneinfo import ZoneInfo
from flask import Flask, render_template_string, request, redirect, send_file
import sqlite3
from io import BytesIO
import re

app = Flask(__name__)
DATABASE_URL = os.environ.get('DATABASE_URL')
USE_POSTGRES = False
if DATABASE_URL:
    try:
        import psycopg2
        USE_POSTGRES = True
    except:
        USE_POSTGRES = False

def get_conn():
    if USE_POSTGRES:
        return psycopg2.connect(DATABASE_URL, sslmode='require')
    else:
        conn = sqlite3.connect('oficina.db')
        conn.row_factory = sqlite3.Row
        return conn

def init_db():
    conn = get_conn()
    cur = conn.cursor()
    if USE_POSTGRES:
        cur.execute('CREATE TABLE IF NOT EXISTS produtos (id SERIAL PRIMARY KEY, nome TEXT, preco REAL, estoque INTEGER)')
        cur.execute('CREATE TABLE IF NOT EXISTS servicos (id SERIAL PRIMARY KEY, cliente TEXT, carro TEXT, descricao TEXT, valor REAL, data TEXT, placa TEXT, veiculo TEXT, produtos_usados TEXT)')
        for col in ["cliente","carro","placa","veiculo","descricao","valor","data","produtos_usados"]:
            try: cur.execute(f'ALTER TABLE servicos ADD COLUMN IF NOT EXISTS {col} TEXT')
            except: pass
    else:
        cur.execute('CREATE TABLE IF NOT EXISTS produtos (id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT, preco REAL, estoque INTEGER)')
        cur.execute('CREATE TABLE IF NOT EXISTS servicos (id INTEGER PRIMARY KEY AUTOINCREMENT, cliente TEXT, placa TEXT, veiculo TEXT, carro TEXT, descricao TEXT, valor REAL, data TEXT, produtos_usados TEXT)')
    conn.commit()
    conn.close()

def buscar_servicos(filtro=""):
    conn = get_conn()
    cur = conn.cursor()
    if filtro:
        if USE_POSTGRES:
            cur.execute("SELECT id, cliente, placa, veiculo, carro, descricao, valor, data, produtos_usados FROM servicos WHERE cliente ILIKE %s OR placa ILIKE %s OR veiculo ILIKE %s OR descricao ILIKE %s ORDER BY id DESC", (f"%{filtro}%",f"%{filtro}%",f"%{filtro}%",f"%{filtro}%"))
        else:
            cur.execute("SELECT id, cliente, placa, veiculo, carro, descricao, valor, data, produtos_usados FROM servicos WHERE cliente LIKE? OR placa LIKE? OR veiculo LIKE? OR descricao LIKE? ORDER BY id DESC", (f"%{filtro}%",f"%{filtro}%",f"%{filtro}%",f"%{filtro}%"))
    else:
        cur.execute("SELECT id, cliente, placa, veiculo, carro, descricao, valor, data, produtos_usados FROM servicos ORDER BY id DESC")
    rows = cur.fetchall()
    conn.close()
    lista=[]
    for r in rows:
        if USE_POSTGRES:
            id_, cliente, placa, veiculo, carro, descricao, valor, data, prod = r
        else:
            id_=r['id']; cliente=r['cliente']; placa=r['placa']; veiculo=r['veiculo']; carro=r['carro']; descricao=r['descricao']; valor=r['valor']; data=r['data']; prod=r['produtos_usados']
        lista.append({"id":id_, "cliente":cliente or "", "placa":placa or "", "veiculo":veiculo or carro or "", "descricao":descricao or "", "valor":valor or 0, "data":data or "", "prod":prod or ""})
    return lista

BASE = """
<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Oficina PRO</title>
<link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/css/bootstrap.min.css" rel="stylesheet">
<style>
body{background:#0e0e0e;color:#fff;font-family:Inter,Arial}
.navbar{background:#000!important;border-bottom:3px solid #f7b500;padding:12px 15px}
.card-dark{background:#1c1c1c;border:1px solid #2a2a2a;border-radius:12px}
.btn-yellow{background:#f7b500;color:#000;font-weight:800;border-radius:8px;border:none}
.btn-dark2{background:#2a2a2a;color:#fff;border:none;border-radius:8px}
.table{--bs-table-bg:#1c1c1c!important;--bs-table-color:#e0e0e0!important;background:#1c1c1c!important;margin:0}
.table > :not(caption) > * > *{background-color:#1c1c1c!important;color:#ddd!important;box-shadow:none!important;border-color:#222!important}
.table thead th{background:#1c1c1c!important;color:#fff!important;border-bottom:1px solid #333!important}
.table tbody td{background:#1c1c1c!important;color:#ddd!important;border-bottom:1px solid #222!important}
.small-label{color:#888;font-size:13px}
.big-number{font-size:22px;font-weight:700}
.produto-item{background:#252525;border:1px solid #333;border-radius:8px;padding:10px 12px;margin-bottom:8px;display:flex;justify-content:space-between;align-items:center;cursor:pointer}
.produto-item.selected{background:#2e2a15;border-color:#f7b500}
.qty-box{display:none; align-items:center; gap:6px;}
.qty-box input{width:55px; text-align:center; background:#000; border:1px solid #f7b500; color:#f7b500; font-weight:900; border-radius:6px; padding:4px;}
.qty-btn{background:#f7b500; color:#000; border:none; width:28px; height:28px; border-radius:6px; font-weight:900;}
</style>
</head><body>
<nav class="navbar d-flex justify-content-between">
<div class="fw-bold">OFICINA PRO - NA NUVEM</div>
<div class="d-flex gap-2 flex-wrap">
<a href="/" class="btn btn-dark2 btn-sm">Início</a>
<a href="/produtos" class="btn btn-dark2 btn-sm">Produtos</a>
<a href="/historico" class="btn btn-dark2 btn-sm">Histórico</a>
<a href="/exportar_excel" class="btn btn-dark2 btn-sm">Excel OS</a>
<a href="/exportar_produtos" class="btn btn-dark2 btn-sm">Excel Produtos</a>
<a href="/nova_os" class="btn btn-yellow btn-sm">+ Nova OS</a>
</div>
</nav>
<div class="container py-4" style="max-width:950px">{{content|safe}}</div>
</body></html>
"""

@app.route('/')
def index():
    init_db()
    conn = get_conn(); cur = conn.cursor()
    try: cur.execute("SELECT COUNT(*) FROM servicos"); total_serv = cur.fetchone()[0]
    except: total_serv=0
    try: cur.execute("SELECT COALESCE(SUM(valor),0) FROM servicos"); faturado = cur.fetchone()[0] or 0
    except: faturado=0
    try: cur.execute("SELECT COUNT(*) FROM produtos"); total_prod = cur.fetchone()[0]
    except: total_prod=0
    conn.close()
    servicos = buscar_servicos()[:8]
    linhas=""
    for s in servicos:
        linhas+=f"<tr><td>{s['data']}</td><td><b>{s['cliente']}</b><br><span class='small-label'>{s['veiculo']} {s['placa']}</span></td><td>{s['descricao']}<br><span class='small-label' style='color:#f7b500'>{s['prod']}</span></td><td>R$ {float(s['valor']):.2f}</td><td><a href='/imprimir/{s['id']}' target='_blank' class='btn btn-sm btn-dark2'>🖨️</a> <a href='/editar_servico/{s['id']}' class='btn btn-sm btn-dark2'>✏️</a> <a href='/excluir_servico/{s['id']}' class='btn btn-sm btn-danger'>🗑️</a></td></tr>"
    html = f"""
    <div class="row g-3">
      <div class="col-6"><div class="card-dark p-4"><div class="big-number">{total_serv}</div><div class="small-label mt-2">Serviços</div></div></div>
      <div class="col-6"><div class="card-dark p-4"><div class="big-number">R$ {float(faturado):.2f}</div><div class="small-label mt-2">Faturado</div></div></div>
      <div class="col-6"><div class="card-dark p-4"><div class="big-number">{total_prod}</div><div class="small-label mt-2">Produtos</div></div></div>
      <div class="col-6"><div class="card-dark p-4 d-flex flex-column gap-2"><a href="/nova_os" class="btn btn-yellow w-100 py-3">+ REGISTRAR SERVIÇO</a><div class="d-flex gap-2"><a href="/exportar_excel" class="btn btn-dark2 w-50">📊 Excel OS</a><a href="/exportar_produtos" class="btn btn-dark2 w-50">📦 Excel Produtos</a></div></div></div>
    </div>
    <div class="card-dark p-4 mt-4">
      <h6 class="fw-bold mb-3">Últimos Serviços</h6>
      <table class="table"><thead><tr><th>Data</th><th>Cliente</th><th>Feito</th><th>Valor</th><th>Ações</th></tr></thead><tbody>{linhas if linhas else '<tr><td colspan=5 class=small-label>Nenhum serviço</td></tr>'}</tbody></table>
    </div>
    """
    return render_template_string(BASE, content=html)

@app.route('/produtos', methods=['GET','POST'])
def produtos():
    init_db()
    conn = get_conn(); cur = conn.cursor()
    msg = ""
    if request.method == 'POST':
        if 'arquivo' in request.files and request.files['arquivo'].filename!= '':
            file = request.files['arquivo']
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file)
                sheet = None
                for name in wb.sheetnames:
                    ws_tmp = wb[name]
                    headers_tmp = [str(ws_tmp.cell(1, c).value or '').upper() for c in range(1, 10)]
                    if any('PRODUTO' in h or 'NOME' in h for h in headers_tmp):
                        sheet = ws_tmp
                        break
                if sheet is None:
                    sheet = wb.active

                headers = {}
                for c in range(1, sheet.max_column+1):
                    val = str(sheet.cell(1, c).value or '').lower().strip()
                    headers[val] = c

                def find_col(names):
                    for n in names:
                        for h, col in headers.items():
                            if n in h:
                                return col
                    return None

                c_prod = find_col(['produto','nome','descricao','descrição','item'])
                c_preco = find_col(['valor revenda','revenda','preco','preço','valor','price'])
                c_est = find_col(['estoque','qtd','quantidade','colunas1','saldo','est'])

                if not c_prod: c_prod = 1
                if not c_preco: c_preco = 2

                count = 0
                for r in range(2, sheet.max_row+1):
                    nome = str(sheet.cell(r, c_prod).value or '').strip()
                    if not nome or nome.lower() in ['none','nan',''] or 'total' in nome.lower():
                        continue
                    try:
                        preco = float(sheet.cell(r, c_preco).value or 0)
                    except:
                        preco = 0.0
                    try:
                        estoque = int(float(sheet.cell(r, c_est).value or 0)) if c_est else 0
                    except:
                        estoque = 0
                    if preco <= 0 and estoque <= 0:
                        continue

                    if USE_POSTGRES:
                        cur.execute("SELECT id FROM produtos WHERE UPPER(nome)=UPPER(%s)", (nome,))
                    else:
                        cur.execute("SELECT id FROM produtos WHERE UPPER(nome)=UPPER(?)", (nome,))
                    existe = cur.fetchone()
                    if existe:
                        if USE_POSTGRES:
                            cur.execute("UPDATE produtos SET preco=%s, estoque=%s WHERE UPPER(nome)=UPPER(%s)", (preco, estoque, nome))
                        else:
                            cur.execute("UPDATE produtos SET preco=?, estoque=? WHERE UPPER(nome)=UPPER(?)", (preco, estoque, nome))
                    else:
                        if USE_POSTGRES:
                            cur.execute("INSERT INTO produtos (nome, preco, estoque) VALUES (%s,%s,%s)", (nome, preco, estoque))
                        else:
                            cur.execute("INSERT INTO produtos (nome, preco, estoque) VALUES (?,?,?)", (nome, preco, estoque))
                        count+=1
                conn.commit()
                msg = f"<div class='alert alert-success'>✅ {count} produtos novos importados! {sheet.max_row-1} linhas lidas. Produtos existentes foram atualizados.</div>"
            except Exception as e:
                msg = f"<div class='alert alert-danger'>❌ Erro ao importar: {e}</div>"
        else:
            if USE_POSTGRES: cur.execute("INSERT INTO produtos (nome, preco, estoque) VALUES (%s,%s,%s)", (request.form['nome'], request.form['preco'], request.form['estoque']))
            else: cur.execute("INSERT INTO produtos (nome, preco, estoque) VALUES (?,?,?)", (request.form['nome'], request.form['preco'], request.form['estoque']))
            conn.commit(); conn.close(); return redirect('/produtos')

    cur.execute("SELECT * FROM produtos ORDER BY id DESC"); rows = cur.fetchall(); conn.close()
    linhas=""
    for r in rows:
        id_, nome, preco, est = (r[0], r[1], r[2], r[3]) if USE_POSTGRES else (r['id'], r['nome'], r['preco'], r['estoque'])
        linhas+=f"<tr><td>{nome}</td><td>R$ {float(preco or 0):.2f}</td><td>{est}</td><td><a href='/editar_produto/{id_}' class='btn btn-sm btn-dark2' style='margin-right:4px;'>✏️</a><a href='/excluir_produto/{id_}' class='btn btn-sm btn-danger'>X</a></td></tr>"
    html = f"""
    {msg}
    <div class="card-dark p-4">
    <h5 class="fw-bold mb-3">Produtos ({len(rows)})</h5>
    <form method="POST" class="row g-2 mb-4">
      <div class="col-md-5"><input name="nome" class="form-control bg-dark text-white border-secondary" placeholder="Nome" required></div>
      <div class="col-md-3"><input name="preco" type="number" step="0.01" class="form-control bg-dark text-white border-secondary" placeholder="Preço" required></div>
      <div class="col-md-2"><input name="estoque" type="number" class="form-control bg-dark text-white border-secondary" placeholder="Est" required></div>
      <div class="col-md-2"><button class="btn btn-yellow w-100">+ Add</button></div>
    </form>

    <div class="row g-2 mb-4">
      <div class="col-md-6">
        <div style="background:#1a1a00; border:2px dashed #f7b500; padding:12px; border-radius:10px">
          <h6 class="fw-bold" style="color:#f7b500">📥 Importar Planilha</h6>
          <p class="small-label">Use sua PLANILHA_OFICINA_CERTA.xlsx - lê PRODUTO / VALOR REVENDA / ESTOQUE automaticamente</p>
          <form method="POST" enctype="multipart/form-data" class="d-flex gap-2">
            <input type="file" name="arquivo" accept=".xlsx,.xls" class="form-control bg-dark text-white border-secondary" required>
            <button class="btn btn-yellow">IMPORTAR</button>
          </form>
        </div>
      </div>
      <div class="col-md-6">
        <div style="background:#0a1a0a; border:2px dashed #22c55e; padding:12px; border-radius:10px">
          <h6 class="fw-bold" style="color:#22c55e">📤 Exportar Planilha</h6>
          <p class="small-label">Baixa todos os {len(rows)} produtos em Excel para backup</p>
          <a href="/exportar_produtos" class="btn w-100" style="background:#22c55e; color:black; font-weight:800">📊 BAIXAR EXCEL PRODUTOS</a>
        </div>
      </div>
    </div>

    <table class="table"><thead><tr><th>Nome</th><th>Preço</th><th>Est</th><th>Ações</th></tr></thead><tbody>{linhas if linhas else '<tr><td colspan=4 class=small-label>Nenhum produto ainda</td></tr>'}</tbody></table>
    </div>
    """
    return render_template_string(BASE, content=html)

@app.route('/editar_produto/<int:id>', methods=['GET','POST'])
def editar_produto(id):
    init_db()
    conn = get_conn(); cur = conn.cursor()
    cur.execute("SELECT * FROM produtos WHERE id=%s" % ("%s" if USE_POSTGRES else "?"), (id,))
    r = cur.fetchone()
    if not r: conn.close(); return redirect('/produtos')
    if USE_POSTGRES: pid, nome, preco, est = r
    else: pid=r['id']; nome=r['nome']; preco=r['preco']; est=r['estoque']
    if request.method=='POST':
        if USE_POSTGRES: cur.execute("UPDATE produtos SET nome=%s, preco=%s, estoque=%s WHERE id=%s", (request.form['nome'], request.form['preco'], request.form['estoque'], id))
        else: cur.execute("UPDATE produtos SET nome=?, preco=?, estoque=? WHERE id=?", (request.form['nome'], request.form['preco'], request.form['estoque'], id))
        conn.commit(); conn.close(); return redirect('/produtos')
    conn.close()
    html=f"""
    <div class="card-dark p-4">
      <h5 class="fw-bold mb-3">Editar Produto #{pid}</h5>
      <form method="POST" class="row g-3">
        <div class="col-md-6"><label class="small-label">Nome</label><input name="nome" value="{nome}" class="form-control bg-dark text-white border-secondary" required></div>
        <div class="col-md-3"><label class="small-label">Preço</label><input name="preco" type="number" step="0.01" value="{preco}" class="form-control bg-dark text-white border-secondary" required></div>
        <div class="col-md-3"><label class="small-label">Estoque</label><input name="estoque" type="number" value="{est}" class="form-control bg-dark text-white border-secondary" required></div>
        <div class="col-12"><button class="btn btn-yellow w-100">SALVAR EDIÇÃO</button><a href="/produtos" class="btn btn-dark2 w-100 mt-2">Cancelar</a></div>
      </form>
    </div>
    """
    return render_template_string(BASE, content=html)

@app.route('/excluir_produto/<int:id>')
def excluir_produto(id):
    conn = get_conn(); cur = conn.cursor()
    if USE_POSTGRES: cur.execute("DELETE FROM produtos WHERE id=%s", (id,))
    else: cur.execute("DELETE FROM produtos WHERE id=?", (id,))
    conn.commit(); conn.close(); return redirect('/produtos')

@app.route('/excluir_servico/<int:id>')
def excluir_servico(id):
    conn = get_conn(); cur = conn.cursor()
    if USE_POSTGRES: cur.execute("DELETE FROM servicos WHERE id=%s", (id,))
    else: cur.execute("DELETE FROM servicos WHERE id=?", (id,))
    conn.commit(); conn.close(); return redirect(request.referrer or '/')

@app.route('/nova_os', methods=['GET','POST'])
def nova_os():
    init_db()
    conn = get_conn(); cur = conn.cursor()
    cur.execute("SELECT * FROM produtos WHERE estoque > 0 ORDER BY nome ASC"); produtos = cur.fetchall()
    if request.method == 'POST':
        cliente = request.form.get('cliente',''); placa = request.form.get('placa',''); veiculo = request.form.get('veiculo','')
        descricao = request.form.get('descricao',''); valor = request.form.get('valor_total','0') or 0
        selecionados = request.form.getlist('produtos')
        data = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y %H:%M")
        nomes_usados=[]
        for pid in selecionados:
            try:
                qty = int(request.form.get(f'qty_{pid}', '1') or 1)
                if qty < 1: qty=1
                if USE_POSTGRES: cur.execute("SELECT nome FROM produtos WHERE id=%s", (int(pid),))
                else: cur.execute("SELECT nome FROM produtos WHERE id=?", (int(pid),))
                prow = cur.fetchone()
                if prow:
                    nome_p = prow[0] if USE_POSTGRES else prow['nome']
                    nomes_usados.append(f"{qty}x {nome_p}" if qty>1 else nome_p)
                    if USE_POSTGRES: cur.execute("UPDATE produtos SET estoque = estoque - %s WHERE id=%s", (qty, int(pid),))
                    else: cur.execute("UPDATE produtos SET estoque = estoque -? WHERE id=?", (qty, int(pid),))
            except Exception as e: print(e); pass
        prod_txt = ", ".join(nomes_usados)
        carro_full = f"{veiculo} - {placa}".strip(" -")
        if USE_POSTGRES: cur.execute("INSERT INTO servicos (cliente, carro, placa, veiculo, descricao, valor, data, produtos_usados) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)", (cliente, carro_full, placa, veiculo, descricao, valor, data, prod_txt))
        else: cur.execute("INSERT INTO servicos (cliente, placa, veiculo, carro, descricao, valor, data, produtos_usados) VALUES (?,?,?,?,?,?,?,?)", (cliente, placa, veiculo, carro_full, descricao, valor, data, prod_txt))
        conn.commit(); conn.close(); return redirect('/')
    lista_prod_html=""
    for r in produtos:
        id_, nome, preco, est = (r[0], r[1], r[2], r[3]) if USE_POSTGRES else (r['id'], r['nome'], r['preco'], r['estoque'])
        lista_prod_html += f"""
        <div class="produto-item" id="item-{id_}" data-preco="{float(preco or 0)}" onclick="toggleProd({id_})">
          <div><input type="checkbox" name="produtos" value="{id_}" id="chk-{id_}" style="display:none"><b>{nome}</b><br><span class="small-label">R$ {float(preco or 0):.2f} - Est: {est}</span></div>
          <div style="display:flex; align-items:center; gap:8px;">
            <div class="qty-box" id="qtybox-{id_}" onclick="event.stopPropagation()">
              <button type="button" class="qty-btn" onclick="changeQty({id_},-1)">-</button>
              <input type="number" id="qty-{id_}" name="qty_{id_}" value="1" min="1" oninput="calcTotal()" onclick="event.stopPropagation()">
              <button type="button" class="qty-btn" onclick="changeQty({id_},1)">+</button>
            </div>
            <span class="small-label">✓</span>
          </div>
        </div>
        """
    if not lista_prod_html: lista_prod_html = "<div class='small-label'>Nenhum produto em estoque.</div>"
    conn.close()
    html = f"""
    <div class="card-dark p-4">
    <h5 class="fw-bold mb-3">+ Nova OS</h5>
    <form method="POST">
      <div class="row g-3">
        <div class="col-md-6"><label class="small-label">Cliente *</label><input name="cliente" class="form-control bg-dark text-white border-secondary" required></div>
        <div class="col-md-3"><label class="small-label">Placa</label><input name="placa" class="form-control bg-dark text-white border-secondary"></div>
        <div class="col-md-3"><label class="small-label">Veículo</label><input name="veiculo" class="form-control bg-dark text-white border-secondary"></div>
        <div class="col-12"><label class="small-label">O que foi feito *</label><textarea name="descricao" rows="3" class="form-control bg-dark text-white border-secondary" required></textarea></div>
        <div class="col-md-4"><label class="small-label">Mão de Obra R$</label><input id="valor_mao" type="number" step="0.01" class="form-control bg-dark text-white border-secondary" placeholder="0.00" oninput="calcTotal()"></div>
        <div class="col-md-4">
          <div class="d-flex justify-content-between align-items-center">
            <label class="small-label">Total Automático</label>
            <button type="button" id="btnEditarNova" onclick="toggleEditarNova()" class="btn btn-sm py-0 px-2" style="background:#f7b500;color:#000;font-weight:800;font-size:11px">✏️ Editar</button>
          </div>
          <input id="valor_total" name="valor_total" type="number" step="0.01" class="form-control bg-dark text-white border-secondary" style="border-color:#f7b500!important;font-weight:bold" readonly>
          <div id="valor_display" class="small-label mt-1" style="color:#f7b500">R$ 0.00 - Automático</div>
        </div>
      </div>
      <div class="mt-4"><label class="small-label fw-bold">Produtos usados</label><div class="mt-2" style="max-height:380px;overflow-y:auto">{lista_prod_html}</div></div>
      <button class="btn btn-yellow w-100 py-3 mt-4">SALVAR SERVIÇO - <span id="valor_display2">R$ 0.00</span></button>
    </form>
    </div>
    <script>
    let modoManualNova = false;
    function toggleEditarNova(){{
      modoManualNova =!modoManualNova;
      const input = document.getElementById('valor_total');
      const btn = document.getElementById('btnEditarNova');
      const disp = document.getElementById('valor_display');
      if(modoManualNova){{
        input.removeAttribute('readonly'); input.style.borderColor='#22c55e'; input.style.color='#22c55e'; input.focus();
        btn.innerText='🔒 Auto'; btn.style.background='#22c55e';
        disp.innerText='✏️ Editando manualmente'; disp.style.color='#22c55e';
      }} else {{
        input.setAttribute('readonly','readonly'); input.style.borderColor='#f7b500'; input.style.color='#fff';
        btn.innerText='✏️ Editar'; btn.style.background='#f7b500';
        calcTotal();
      }}
    }}
    function toggleProd(id){{
      const el=document.getElementById('item-'+id);
      const chk=document.getElementById('chk-'+id);
      const qtybox=document.getElementById('qtybox-'+id);
      el.classList.toggle('selected');
      chk.checked = el.classList.contains('selected');
      qtybox.style.display = chk.checked? 'flex' : 'none';
      if(!chk.checked){{ document.getElementById('qty-'+id).value=1; }}
      calcTotal();
    }}
    function changeQty(id, delta){{
      const inp=document.getElementById('qty-'+id);
      let v=parseInt(inp.value)||1;
      v+=delta;
      if(v<1) v=1;
      inp.value=v;
      calcTotal();
    }}
    function calcTotal(){{
      if(modoManualNova) return;
      let t=parseFloat(document.getElementById('valor_mao').value)||0;
      document.querySelectorAll('.produto-item.selected').forEach(item=>{{
        const id=item.id.replace('item-','');
        const preco=parseFloat(item.dataset.preco)||0;
        const qty=parseInt(document.getElementById('qty-'+id).value)||1;
        t+=preco*qty;
      }});
      document.getElementById('valor_total').value=t.toFixed(2);
      document.getElementById('valor_display').innerText='R$ '+t.toFixed(2)+' - Automático';
      const d2=document.getElementById('valor_display2');
      if(d2) d2.innerText='R$ '+t.toFixed(2);
    }}
    </script>
    """
    return render_template_string(BASE, content=html)

@app.route('/editar_servico/<int:id>', methods=['GET','POST'])
def editar_servico(id):
    init_db()
    conn = get_conn(); cur = conn.cursor()
    if USE_POSTGRES: cur.execute("SELECT id, cliente, placa, veiculo, carro, descricao, valor, data, produtos_usados FROM servicos WHERE id=%s", (id,))
    else: cur.execute("SELECT id, cliente, placa, veiculo, carro, descricao, valor, data, produtos_usados FROM servicos WHERE id=?", (id,))
    r = cur.fetchone()
    if not r: conn.close(); return redirect('/')
    if USE_POSTGRES:
        s_id, s_cliente, s_placa, s_veiculo, s_carro, s_desc, s_valor, s_data, s_prod = r
    else:
        s_id=r['id']; s_cliente=r['cliente']; s_placa=r['placa']; s_veiculo=r['veiculo']; s_carro=r['carro']; s_desc=r['descricao']; s_valor=r['valor']; s_data=r['data']; s_prod=r['produtos_usados']
    produtos_existentes = {}
    if s_prod:
        for part in s_prod.split(','):
            part=part.strip()
            if not part: continue
            m = re.match(r'(\d+)x\s+(.+)', part, re.I)
            if m:
                produtos_existentes[m.group(2).strip().upper()] = int(m.group(1))
            else:
                produtos_existentes[part.strip().upper()] = 1
    if request.method == 'POST':
        cliente = request.form.get('cliente',''); placa = request.form.get('placa',''); veiculo = request.form.get('veiculo','')
        descricao = request.form.get('descricao',''); valor = request.form.get('valor_total') or request.form.get('valor') or 0
        selecionados = request.form.getlist('produtos')
        nomes_usados=[]
        for pid in selecionados:
            try:
                qty = int(request.form.get(f'qty_{pid}', '1') or 1)
                if USE_POSTGRES: cur.execute("SELECT nome FROM produtos WHERE id=%s", (int(pid),))
                else: cur.execute("SELECT nome FROM produtos WHERE id=?", (int(pid),))
                prow = cur.fetchone()
                if prow:
                    nome_p = prow[0] if USE_POSTGRES else prow['nome']
                    nomes_usados.append(f"{qty}x {nome_p}" if qty>1 else nome_p)
            except: pass
        prod_txt = ", ".join(nomes_usados)
        if USE_POSTGRES: cur.execute("UPDATE servicos SET cliente=%s, placa=%s, veiculo=%s, descricao=%s, valor=%s, produtos_usados=%s WHERE id=%s", (cliente, placa, veiculo, descricao, valor, prod_txt, id))
        else: cur.execute("UPDATE servicos SET cliente=?, placa=?, veiculo=?, descricao=?, valor=?, produtos_usados=? WHERE id=?", (cliente, placa, veiculo, descricao, valor, prod_txt, id))
        conn.commit(); conn.close(); return redirect('/historico')
    cur.execute("SELECT * FROM produtos ORDER BY nome ASC")
    produtos_all = cur.fetchall()
    conn.close()
    lista_prod_html=""
    for prod_r in produtos_all:
        pid, nome, preco, est = (prod_r[0], prod_r[1], prod_r[2], prod_r[3]) if USE_POSTGRES else (prod_r['id'], prod_r['nome'], prod_r['preco'], prod_r['estoque'])
        nome_up = nome.strip().upper()
        ja_usado = nome_up in produtos_existentes
        qty_atual = produtos_existentes.get(nome_up, 1)
        selected_class = "selected" if ja_usado else ""
        checked_attr = "checked" if ja_usado else ""
        display_qty = "flex" if ja_usado else "none"
        lista_prod_html += f"""
        <div class="produto-item {selected_class}" id="item-{pid}" data-preco="{float(preco or 0)}" onclick="toggleProd({pid})">
          <div><input type="checkbox" name="produtos" value="{pid}" id="chk-{pid}" style="display:none" {checked_attr}><b>{nome}</b><br><span class="small-label">R$ {float(preco or 0):.2f} - Est: {est}</span></div>
          <div style="display:flex; align-items:center; gap:8px;">
            <div class="qty-box" id="qtybox-{pid}" style="display:{display_qty}" onclick="event.stopPropagation()">
              <button type="button" class="qty-btn" onclick="changeQty({pid},-1)">-</button>
              <input type="number" id="qty-{pid}" name="qty_{pid}" value="{qty_atual}" min="1" oninput="calcTotal()" onclick="event.stopPropagation()">
              <button type="button" class="qty-btn" onclick="changeQty({pid},1)">+</button>
            </div>
            <span class="small-label">✓</span>
          </div>
        </div>
        """
    html = f"""
    <div class="card-dark p-4">
    <h5 class="fw-bold mb-3">Editar OS #{s_id}</h5>
    <form method="POST">
      <div class="row g-3">
        <div class="col-md-6"><label class="small-label">Cliente</label><input name="cliente" value="{s_cliente}" class="form-control bg-dark text-white border-secondary" required></div>
        <div class="col-md-3"><label class="small-label">Placa</label><input name="placa" value="{s_placa}" class="form-control bg-dark text-white border-secondary"></div>
        <div class="col-md-3"><label class="small-label">Veículo</label><input name="veiculo" value="{s_veiculo or s_carro or ''}" class="form-control bg-dark text-white border-secondary"></div>
        <div class="col-12"><label class="small-label">O que foi feito</label><textarea name="descricao" rows="3" class="form-control bg-dark text-white border-secondary" required>{s_desc}</textarea></div>
        <div class="col-md-4">
          <div class="d-flex justify-content-between align-items-center mb-1">
            <label class="small-label m-0">Total</label>
            <button type="button" id="btnEditar" onclick="toggleEditar()" class="btn btn-sm" style="background:#f7b500;color:#000;font-weight:800;font-size:11px;padding:2px 10px">✏️ Editar</button>
          </div>
          <input id="valor_total" name="valor_total" type="number" step="0.01" value="{s_valor}" class="form-control bg-dark text-white border-secondary" readonly style="border-color:#f7b500!important;font-weight:bold">
          <div id="valor_display" class="small-label mt-1" style="color:#f7b500">R$ {float(s_valor or 0):.2f} - Automático</div>
        </div>
      </div>
      <div class="mt-4"><label class="small-label fw-bold">Produtos usados</label><div class="mt-2" style="max-height:380px;overflow-y:auto">{lista_prod_html}</div></div>
      <button class="btn btn-yellow w-100 py-3 mt-4">SALVAR EDIÇÃO</button>
      <a href="/historico" class="btn btn-dark2 w-100 mt-2">Cancelar</a>
    </form>
    </div>
    <script>
    let modoManual = false;
    function toggleEditar(){{
      modoManual =!modoManual;
      const input = document.getElementById('valor_total');
      const btn = document.getElementById('btnEditar');
      const disp = document.getElementById('valor_display');
      if(modoManual){{
        input.removeAttribute('readonly');
        input.style.borderColor='#22c55e';
        input.style.color='#22c55e';
        input.focus();
        btn.innerText='🔒 Auto';
        btn.style.background='#22c55e';
        disp.innerText='✏️ Editando manualmente - digite o valor';
        disp.style.color='#22c55e';
      }} else {{
        input.setAttribute('readonly','readonly');
        input.style.borderColor='#f7b500';
        input.style.color='#fff';
        btn.innerText='✏️ Editar';
        btn.style.background='#f7b500';
        calcTotal();
      }}
    }}
    function toggleProd(id){{
      const el=document.getElementById('item-'+id);
      const chk=document.getElementById('chk-'+id);
      const qtybox=document.getElementById('qtybox-'+id);
      el.classList.toggle('selected');
      chk.checked = el.classList.contains('selected');
      qtybox.style.display = chk.checked? 'flex' : 'none';
      calcTotal();
    }}
    function changeQty(id, delta){{
      const inp=document.getElementById('qty-'+id);
      let v=parseInt(inp.value)||1; v+=delta; if(v<1) v=1; inp.value=v; calcTotal();
    }}
    function calcTotal(){{
      if(modoManual) return;
      let t=0;
      document.querySelectorAll('.produto-item.selected').forEach(item=>{{
        const id=item.id.replace('item-','');
        const preco=parseFloat(item.dataset.preco)||0;
        const qty=parseInt(document.getElementById('qty-'+id).value)||1;
        t+=preco*qty;
      }});
      document.getElementById('valor_total').value=t.toFixed(2);
      const vd=document.getElementById('valor_display');
      if(vd) vd.innerText='R$ '+t.toFixed(2)+' - Automático';
    }}
    </script>
    """
    return render_template_string(BASE, content=html)

@app.route('/historico')
def historico():
    init_db()
    q = request.args.get('q','')
    servicos = buscar_servicos(q)
    linhas=""
    for s in servicos:
        linhas+=f"<tr><td>{s['data']}</td><td><b>{s['cliente']}</b><br><span class='small-label'>{s['veiculo']} {s['placa']}</span></td><td>{s['descricao']}<br><span class='small-label' style='color:#f7b500'>{s['prod']}</span></td><td>R$ {float(s['valor'] or 0):.2f}</td><td><a href='/imprimir/{s['id']}' target='_blank' class='btn btn-sm btn-dark2'>🖨️</a> <a href='/editar_servico/{s['id']}' class='btn btn-sm btn-dark2'>✏️</a> <a href='/excluir_servico/{s['id']}' class='btn btn-sm btn-danger'>🗑️</a></td></tr>"
    html = f"""
    <div class="card-dark p-4">
    <div class="d-flex justify-content-between align-items-center mb-3"><h5 class="fw-bold m-0">Histórico</h5><div class="d-flex gap-2"><a href="/exportar_excel" class="btn btn-yellow btn-sm">📊 Excel OS</a><a href="/exportar_produtos" class="btn btn-dark2 btn-sm">📦 Excel Produtos</a></div></div>
    <form method="GET" class="mb-3"><div class="input-group"><input name="q" value="{q}" class="form-control bg-dark text-white border-secondary" placeholder="Pesquisar..."><button class="btn btn-yellow">Buscar</button></div></form>
    <table class="table"><thead><tr><th>Data</th><th>Cliente</th><th>Feito</th><th>Valor</th><th>Ações</th></tr></thead><tbody>{linhas if linhas else '<tr><td colspan=5 class=small-label>Nenhum resultado</td></tr>'}</tbody></table>
    </div>
    """
    return render_template_string(BASE, content=html)

@app.route('/exportar_excel')
def exportar_excel():
    servicos = buscar_servicos()
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Serviços"
        ws.append(["Data","Cliente","Placa","Veículo","Descrição","Produtos Usados","Valor"])
        for s in servicos:
            ws.append([s['data'], s['cliente'], s['placa'], s['veiculo'], s['descricao'], s['prod'], float(s['valor'] or 0)])
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(bio, as_attachment=True, download_name=f"oficina_{datetime.now(ZoneInfo('America/Sao_Paulo')).strftime('%d-%m-%Y')}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except:
        output = "Data,Cliente,Placa,Veiculo,Descricao,Produtos,Valor\n"
        for s in servicos:
            output += f"\"{s['data']}\",\"{s['cliente']}\",\"{s['placa']}\",\"{s['veiculo']}\",\"{s['descricao']}\",\"{s['prod']}\",{s['valor']}\n"
        bio = BytesIO(output.encode('utf-8'))
        return send_file(bio, as_attachment=True, download_name="oficina.csv", mimetype="text/csv")

@app.route('/exportar_produtos')
def exportar_produtos():
    init_db()
    conn = get_conn(); cur = conn.cursor()
    cur.execute("SELECT * FROM produtos ORDER BY nome ASC")
    rows = cur.fetchall()
    conn.close()
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(["ID","PRODUTO","VALOR REVENDA","ESTOQUE"])
    for r in rows:
        id_, nome, preco, est = (r[0], r[1], r[2], r[3]) if USE_POSTGRES else (r['id'], r['nome'], r['preco'], r['estoque'])
        ws.append([id_, nome, float(preco or 0), int(est or 0)])
    ws.column_dimensions['B'].width = 45
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=f"produtos_oficina_{datetime.now(ZoneInfo('America/Sao_Paulo')).strftime('%d-%m-%Y')}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route('/imprimir/<int:id>')
def imprimir(id):
    conn = get_conn(); cur = conn.cursor()
    if USE_POSTGRES: cur.execute("SELECT id, cliente, placa, veiculo, carro, descricao, valor, data, produtos_usados FROM servicos WHERE id=%s", (id,))
    else: cur.execute("SELECT id, cliente, placa, veiculo, carro, descricao, valor, data, produtos_usados FROM servicos WHERE id=?", (id,))
    r = cur.fetchone()
    conn.close()
    if not r: return "OS não encontrada"
    if USE_POSTGRES:
        s = {"id":r[0], "cliente":r[1], "placa":r[2], "veiculo":r[3], "descricao":r[5], "valor":r[6], "data":r[7], "prod":r[8]}
    else:
        s = {"id":r['id'], "cliente":r['cliente'], "placa":r['placa'], "veiculo":r['veiculo'], "descricao":r['descricao'], "valor":r['valor'], "data":r['data'], "prod":r['produtos_usados']}
    html_print = f"""
    <!doctype html><html><head><meta charset="utf-8"><title>OS #{s['id']}</title>
    <style>body{{font-family:Arial;padding:30px}}.header{{border-bottom:3px solid #f7b500;padding-bottom:15px;margin-bottom:20px}}.box{{border:1px solid #ddd;padding:15px;border-radius:8px;margin-bottom:15px}}.btn{{background:#f7b500;color:#000;padding:10px 20px;border:none;border-radius:6px;font-weight:bold;cursor:pointer}} @media print{{.no-print{{display:none}}}}</style>
    </head><body>
    <div class="header"><h2>OFICINA PRO - OS #{s['id']}</h2><p>{s['data']}</p></div>
    <div class="box"><b>Cliente:</b> {s['cliente']}<br><b>Veículo:</b> {s['veiculo']}<br><b>Placa:</b> {s['placa']}</div>
    <div class="box"><b>Serviço:</b><br>{s['descricao']}<br><br><b>Produtos:</b> {s['prod']}</div>
    <div class="box"><h3>R$ {float(s['valor'] or 0):.2f}</h3></div>
    <div class="no-print"><button class="btn" onclick="window.print()">🖨️ IMPRIMIR</button></div>
    </body></html>
    """
    return html_print

if __name__ == '__main__':
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 10000)))
